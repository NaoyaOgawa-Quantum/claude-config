#!/usr/bin/env python3
"""short-answer-themes.py — 短答の回答 (文字列の list) をテーマの表 (大きなまとまり → テーマ → 項目 = 表示名と正規表現) で数え、 どの項目にも入らない語の一覧と、 まとまりごとの見出し・人数・抜粋・追記の節を並べた docx を出す。--selftest 内蔵。

授業の事前質問・アンケートの自由記述など、 1 人 1 つの短い回答を「テーマ別に何人」 にまとめる道具。 回ごとに変わるのは
テーマの表 (YAML) だけで、 数え方・語の切り出し・抜粋の取り方・docx の組み方は本 script が持つ。

    python3 short-answer-themes.py --answers answers.json --table themes.yaml            # 人数 + どの項目にも入らない語
    python3 short-answer-themes.py --answers answers.json --table themes.yaml --docx out.docx
    python3 short-answer-themes.py --answers answers.json --table themes.yaml --json     # 機械向けの人数
    python3 short-answer-themes.py --selftest

## 入力

- `--answers` = JSON の list。 要素は回答の文字列 (未回答は "" か null)、 または object (`--field <key>` で回答の欄を指定)。
  **list の長さ = 対象の人数、 空でない要素の数 = 回答した人数**。 氏名・id は入れない (数えるのに要らない)。
- `--table` = YAML (または .json)。 形:

      title: "第 N 回 事前質問の回答まとめ"          # docx の見出し 1 (省略可)
      intro:                                          # 見出し 1 の下の段落 (省略可)
        - "回答 {answers} 人 (対象 {targets} 人)。"
      summary: "大きく分けると次の 3 つ:"             # 省略可。 書けば、 その後にまとまりごとの「名前 — N 人」 の箇条書き
      blocks:                                         # 大きなまとまり → テーマ → 項目
        - name: "Ⅰ 季節"
          themes:
            - name: "暑さ寒さ"
              items:
                - ["暑さ・寒さ", '暑|寒']              # [表示名, 正規表現 (Python re、 flag なし)]
      quotes:                                         # 抜粋 (省略可)
        heading: "目を引いた回答 (原文のまま)"
        keys: ["かき氷", "落ち葉を踏む"]              # 各 key = 回答の部分文字列で、 ちょうど 1 人の回答に当たること
        joiner: "／"                                  # 回答の中の改行を置き換える文字 (既定 ／)
      sections:                                       # 追記の節 (省略可)。 1 要素 = 1 段落
        - h2: "切り口の案"
        - h3: "両方から"
        - b: "「暑さ・寒さ」 は {n:暑さ・寒さ} 人"     # b = 箇条書き / p = 段落 / h2・h3 = 見出し
        - p: "① 太字にする段落"
          bold: true
      font: {name: "Hiragino Kaku Gothic ProN", size: 10.5}   # 本文の字 (省略時はこの値)
      tokenize: {remove: '…', split: '…', strip: '…'}          # 語の切り出しを変えるとき (省略時は下の既定)

  文中の差し込み (title / intro / summary / quotes.heading / sections): `{answers}` = 回答した人数、 `{targets}` = 対象の
  人数、 `{n:<項目の表示名>}` / `{theme:<テーマ名>}` / `{block:<まとまりの名前>}` = その人数。 人数を文に手で写さない
  (表を直すと数が変わる)。 これ以外の `{…}` はそのまま残す。

## 数え方

- 項目の人数 = その正規表現が回答のどこかに当たる人の数 (1 人の中で何度当たっても 1)。
- テーマの人数 = その項目のどれかに当たる人の数 (和集合。 項目の人数の合計ではない)。 まとまりも同じ (テーマの和集合)。
  1 人が複数挙げるので、 テーマ・まとまりの人数の合計は回答した人数を超えうる。
- docx の項目の並び = 人数の多い順 (同数は表の順)、 0 人の項目は出さない (stdout には ⚠️ で出す)。

## どの項目にも入らない語 (表が全員分を拾ったかの確かめ)

回答を語に切り (既定: 括弧 「」｢｣“”" を消し、 読点・句点・改行・中黒・スラッシュ・空白・`*`・行頭の番号 `1.` で区切り、
両端の空白と `-` を落とす)、 **語ごとに**全項目の正規表現を当て、 どれにも当たらない語を回数つきで出す。
表を直す目安 = この一覧が「説明の断片」 (〜だから、 〜な状態、 等) だけになるまで。 語単位で当てるので、 回答全体には
当たる文脈つきの正規表現 (否定の先読みなど) は、 語単位では結果が変わりうる。
あわせて「どの項目にも当たらない回答」 の人数も出す (0 でなければ、 その人の回答はどのテーマにも数えられていない)。

## 抜粋

`quotes.keys` の各 key で回答を 1 つに特定し、 その原文を docx に入れる (手で写さない)。 回答の各行の前後の空白と行頭の
`・` / `* ` を落とし、 行を joiner でつなぐ。 key が 0 人・2 人以上に当たれば exit 2 (人数だけ出し、 回答は出さない)。

## 出力・終了コード

stdout = 人数の表 + どの項目にも入らない語 (`--json` なら同じ内容の JSON)。 `--docx` = docx を書く (既にあれば上書き =
入力から再生成できる生成物)。 共有した文書を上書きする道具ではない (共有先の編集を消さないのは共有側の道具の役目)。
終了コード: 0 = 成功 / 2 = 入力・表の形が想定外 (正規表現の誤り・名前の重複・抜粋の key が 1 人に決まらない・未知の差し込み)。

⚠️ 実際の回答は selftest・例に入れない (selftest は合成の回答だけ)。 docx は python-docx が要る (`--docx` を使う時だけ)。
"""
from __future__ import annotations

import argparse
import collections
import contextlib
import io
import json
import re
import sys
import tempfile
from pathlib import Path

DEFAULT_REMOVE = r'[「」｢｣“”"]'
DEFAULT_SPLIT = r'[、，,。\n・/／]|\s{1,}|\*|(?<=\D)\d+\.'
DEFAULT_STRIP = ' 　-'
DEFAULT_FONT = {"name": "Hiragino Kaku Gothic ProN", "size": 10.5}
DEFAULT_JOINER = "／"
SECTION_KINDS = ("h1", "h2", "h3", "p", "b")
PLACEHOLDER = re.compile(r"\{(answers|targets)\}|\{(n|theme|block):([^{}]+)\}")


class TableError(Exception):
    """入力・表の形が想定外 (exit 2)。"""


# ---------------------------------------------------------------- 読み込み

def load_answers(raw, field=None):
    """JSON の list → (回答の list, 対象の人数)。 回答 = 空でない文字列だけ (元の順)。"""
    if not isinstance(raw, list):
        raise TableError("回答の JSON は list であること")
    out = []
    for i, x in enumerate(raw):
        if isinstance(x, dict):
            if not field:
                raise TableError(f"回答の要素 {i} が object — 回答の欄を --field で指定する")
            x = x.get(field)
        if x is None:
            continue
        if not isinstance(x, str):
            raise TableError(f"回答の要素 {i} が文字列でない ({type(x).__name__})")
        if x.strip():
            out.append(x)
    return out, len(raw)


def load_table(path):
    text = Path(path).read_text(encoding="utf-8")
    if str(path).endswith(".json"):
        return json.loads(text)
    try:
        import yaml
    except ImportError:
        raise TableError("YAML の表を読むには PyYAML が要る (pip install pyyaml)。 .json の表なら不要")
    return yaml.load(text, Loader=getattr(yaml, "CSafeLoader", yaml.SafeLoader))  # safe_load と同じ結果を C 版で


def parse_table(t):
    """表を検査して [(まとまり, [(テーマ, [(表示名, compiled)])])] にする。"""
    if not isinstance(t, dict) or not isinstance(t.get("blocks"), list) or not t["blocks"]:
        raise TableError("表に blocks (まとまりの list) が無い")
    seen = {"block": set(), "theme": set(), "item": set()}

    def uniq(kind, name):
        if not isinstance(name, str) or not name:
            raise TableError(f"{kind} の名前が空か文字列でない: {name!r}")
        if name in seen[kind]:
            raise TableError(f"{kind} の名前が重複: {name}")
        seen[kind].add(name)

    blocks = []
    for b in t["blocks"]:
        if not isinstance(b, dict) or not isinstance(b.get("themes"), list):
            raise TableError(f"まとまりに themes が無い: {b!r}"[:200])
        uniq("block", b.get("name"))
        themes = []
        for th in b["themes"]:
            if not isinstance(th, dict) or not isinstance(th.get("items"), list) or not th["items"]:
                raise TableError(f"テーマに items が無い: {th!r}"[:200])
            uniq("theme", th.get("name"))
            items = []
            for it in th["items"]:
                if not (isinstance(it, (list, tuple)) and len(it) == 2 and all(isinstance(s, str) for s in it)):
                    raise TableError(f"項目は [表示名, 正規表現] の 2 つ組: {it!r} (テーマ {th['name']})")
                label, rx = it
                uniq("item", label)
                try:
                    items.append((label, re.compile(rx)))
                except re.error as e:
                    raise TableError(f"項目 {label} の正規表現が誤り: {e}")
            themes.append((th["name"], items))
        blocks.append((b["name"], themes))
    return blocks


# ---------------------------------------------------------------- 数える

def count(answers, blocks):
    """→ {"items": {表示名: set(回答の番号)}, "themes": {…}, "blocks": {…}}"""
    items, themes, bl = {}, {}, {}
    for bname, ths in blocks:
        bset = set()
        for tname, its in ths:
            tset = set()
            for label, rx in its:
                hit = {i for i, a in enumerate(answers) if rx.search(a)}
                items[label] = hit
                tset |= hit
            themes[tname] = tset
            bset |= tset
        bl[bname] = bset
    return {"items": items, "themes": themes, "blocks": bl}


def tokens(answer, tok=None):
    tok = tok or {}
    a = re.sub(tok.get("remove", DEFAULT_REMOVE), "", answer)
    strip = tok.get("strip", DEFAULT_STRIP)
    out = []
    for p in re.split(tok.get("split", DEFAULT_SPLIT), a):
        p = (p or "").strip(strip)
        if p:
            out.append(p)
    return out


def unmatched_tokens(answers, blocks, tok=None):
    """どの項目の正規表現にも当たらない語 → [(語, 回数)] (多い順)。"""
    rxs = [rx for _, ths in blocks for _, its in ths for _, rx in its]
    c = collections.Counter()
    for a in answers:
        for t in tokens(a, tok):
            if not any(rx.search(t) for rx in rxs):
                c[t] += 1
    return c.most_common()


def fmt_quote(answer, joiner=DEFAULT_JOINER):
    lines = [re.sub(r"^(・|\* )", "", l.strip()) for l in answer.splitlines() if l.strip()]
    return joiner.join(lines)


def pick_quotes(answers, q):
    """quotes.keys → 原文の list。 key が 1 人に決まらなければ TableError (回答の中身は出さない)。"""
    if not q:
        return []
    joiner = q.get("joiner", DEFAULT_JOINER)
    out, bad = [], []
    for k in q.get("keys") or []:
        hit = [a for a in answers if k in a]
        if len(hit) != 1:
            bad.append(f"{k!r} → {len(hit)} 人")
        else:
            out.append(fmt_quote(hit[0], joiner))
    if bad:
        raise TableError("抜粋の key が 1 人の回答に決まらない: " + " / ".join(bad))
    return out


def fill(text, ctx):
    """差し込み ({answers} / {targets} / {n:…} / {theme:…} / {block:…}) を人数に置き換える。"""
    def rep(m):
        if m.group(1):
            return str(ctx[m.group(1)])
        kind, name = m.group(2), m.group(3)
        table = {"n": ctx["counts"]["items"], "theme": ctx["counts"]["themes"], "block": ctx["counts"]["blocks"]}[kind]
        if name not in table:
            raise TableError(f"差し込み {{{kind}:{name}}} の名前が表に無い")
        return str(len(table[name]))
    return PLACEHOLDER.sub(rep, text)


def sorted_items(its, counts):
    """docx に出す項目 = 人数の多い順 (同数は表の順)、 0 人は出さない。"""
    labs = [(label, len(counts["items"][label])) for label, _ in its]
    return [(l, n) for l, n in sorted(labs, key=lambda x: -x[1]) if n > 0]


def parse_sections(raw):
    out = []
    for s in raw or []:
        if not isinstance(s, dict):
            raise TableError(f"sections の要素は mapping: {s!r}"[:200])
        kinds = [k for k in SECTION_KINDS if k in s]
        if len(kinds) != 1 or not isinstance(s[kinds[0]], str):
            raise TableError(f"sections の要素は h1/h2/h3/p/b のどれか 1 つ: {s!r}"[:200])
        extra = set(s) - {kinds[0], "bold"}
        if extra:
            raise TableError(f"sections の要素に知らない key: {sorted(extra)}")
        out.append((kinds[0], s[kinds[0]], bool(s.get("bold"))))
    return out


# ---------------------------------------------------------------- 組み立て

def build(answers_raw, table, field=None):
    """→ 結果の dict (stdout / JSON / docx の共通の源)。"""
    answers, targets = load_answers(answers_raw, field)
    blocks = parse_table(table)
    counts = count(answers, blocks)
    ctx = {"answers": len(answers), "targets": targets, "counts": counts}
    sections = [(k, fill(t, ctx), bold) for k, t, bold in parse_sections(table.get("sections"))]
    q = table.get("quotes") or {}
    res = {
        "answers": len(answers),
        "targets": targets,
        "title": fill(table["title"], ctx) if table.get("title") else None,
        "intro": [fill(p, ctx) for p in table.get("intro") or []],
        "summary": fill(table["summary"], ctx) if table.get("summary") else None,
        "blocks": [
            {"name": b, "n": len(counts["blocks"][b]), "themes": [
                {"name": th, "n": len(counts["themes"][th]),
                 "items": [{"label": l, "n": len(counts["items"][l])} for l, _ in its],
                 "shown": sorted_items(its, counts)}
                for th, its in ths]}
            for b, ths in blocks],
        "quotes_heading": fill(q["heading"], ctx) if q.get("heading") else None,
        "quotes": pick_quotes(answers, q),
        "sections": sections,
        "unmatched_tokens": unmatched_tokens(answers, blocks, table.get("tokenize")),
        "unmatched_answers": sum(1 for i in range(len(answers))
                                 if not any(i in s for s in counts["blocks"].values())),
        "font": {**DEFAULT_FONT, **(table.get("font") or {})},
    }
    return res


def report(res):
    lines = [f"回答 {res['answers']} 人 (対象 {res['targets']} 人)"]
    zero = []
    for b in res["blocks"]:
        lines.append(f"{b['name']} {b['n']}")
        for th in b["themes"]:
            lines.append(f"   {th['name']} {th['n']} " + "、".join(f"{l}({n})" for l, n in th["shown"]))
            zero += [f"{th['name']} / {it['label']}" for it in th["items"] if it["n"] == 0]
    for z in zero:
        lines.append(f"⚠️ 0 人の項目 (docx に出ない): {z}")
    lines.append(f"どの項目にも当たらない回答: {res['unmatched_answers']} 人")
    lines.append(f"--- どの項目にも入らない語 ({len(res['unmatched_tokens'])} 種類)")
    lines += [f"{n} {t}" for t, n in res["unmatched_tokens"]]
    return "\n".join(lines)


def write_docx(res, out):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    st = doc.styles["Normal"]
    if res["font"].get("name"):
        st.font.name = res["font"]["name"]
        st.element.rPr.rFonts.set(qn("w:eastAsia"), res["font"]["name"])
    if res["font"].get("size"):
        st.font.size = Pt(res["font"]["size"])

    if res["title"]:
        doc.add_heading(res["title"], level=1)
    for p in res["intro"]:
        doc.add_paragraph(p)
    if res["summary"]:
        doc.add_paragraph(res["summary"])
        for b in res["blocks"]:
            doc.add_paragraph(f"{b['name']} — {b['n']} 人", style="List Bullet")
    for b in res["blocks"]:
        doc.add_heading(f"{b['name']}（{b['n']} 人）", level=2)
        for th in b["themes"]:
            doc.add_heading(f"{th['name']}（{th['n']} 人）", level=3)
            doc.add_paragraph("、".join(f"{l} ({n})" for l, n in th["shown"]))
    if res["quotes"]:
        if res["quotes_heading"]:
            doc.add_heading(res["quotes_heading"], level=2)
        for q in res["quotes"]:
            doc.add_paragraph(q, style="List Bullet")
    for kind, text, bold in res["sections"]:
        if kind[0] == "h":
            doc.add_heading(text, level=int(kind[1]))
        elif kind == "b":
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(text).bold = bold or None
        else:
            p = doc.add_paragraph()
            p.add_run(text).bold = bold or None
    doc.save(out)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--answers", help="回答の JSON (list)")
    ap.add_argument("--table", help="テーマの表 (YAML / .json)")
    ap.add_argument("--field", help="回答の要素が object のとき、 回答の欄の key")
    ap.add_argument("--docx", help="docx の書き先 (既にあれば上書き)")
    ap.add_argument("--json", action="store_true", help="人数を JSON で出す")
    ap.add_argument("--selftest", action="store_true")
    a = ap.parse_args(argv)
    if a.selftest:
        return selftest()
    if not (a.answers and a.table):
        ap.error("--answers と --table が要る")
    try:
        raw = json.loads(Path(a.answers).read_text(encoding="utf-8"))
        res = build(raw, load_table(a.table), a.field)
    except TableError as e:
        print(f"❌ {e}", file=sys.stderr)
        return 2
    if a.json:
        print(json.dumps({k: res[k] for k in ("answers", "targets", "blocks", "unmatched_tokens", "unmatched_answers")},
                         ensure_ascii=False, indent=1))
    else:
        print(report(res))
    if a.docx:
        write_docx(res, a.docx)
        print(f"書いた: {a.docx}", file=sys.stderr)
    return 0


# ---------------------------------------------------------------- selftest (合成の回答だけ)

ANS = [
    "かき氷、海",                       # 0 夏: 食べ物 + 水辺
    "落ち葉を踏む音\n・紅葉",            # 1 秋: 音 + 色
    "海で泳ぐ、海の家でかき氷",          # 2 夏: 水辺 (1 人の中で 2 回) + 食べ物
    "暑さ",                              # 3 夏: 暑さ
    "",                                  # 未回答
    None,                                # 未回答
    "雪、こたつ、なんとなく好きだから",  # 4 冬 + 説明の断片
    "1.花見\n2.新学期\n3.どれも好き",  # 5 春 (番号つき) + 説明の断片
]
TABLE = {
    "title": "合成の季節アンケート",
    "intro": ["回答 {answers} 人 (対象 {targets} 人)。 かき氷 {n:かき氷} 人、 夏 {block:夏} 人、 {未知}"],
    "summary": "大きく分けると:",
    "blocks": [
        {"name": "夏", "themes": [
            {"name": "水辺と暑さ", "items": [["暑さ", "暑"], ["海", "海"]]},
            {"name": "食べ物", "items": [["かき氷", "かき氷"], ["すいか", "すいか"]]}]},
        {"name": "秋冬春", "themes": [
            {"name": "秋", "items": [["落ち葉・紅葉", "落ち葉|紅葉"], ["音", "音"]]},
            {"name": "冬", "items": [["雪", "雪"], ["こたつ", "こたつ"]]},
            {"name": "春", "items": [["花見", "花見"], ["新学期", "新学期"]]}]},
    ],
    "quotes": {"heading": "抜粋 ({answers} 人から)", "keys": ["落ち葉を踏む", "泳ぐ"]},
    "sections": [{"h2": "案"}, {"p": "① 太字", "bold": True}, {"b": "海は {n:海} 人"}],
}


def selftest() -> int:
    fails = []

    def check(cond, name):
        print(("PASS " if cond else "FAIL ") + name)
        if not cond:
            fails.append(name)

    def err(fn):
        try:
            fn()
        except TableError as e:
            return str(e)
        return ""

    res = build(ANS, TABLE)
    by = {th["name"]: th for b in res["blocks"] for th in b["themes"]}
    items = {it["label"]: it["n"] for th in by.values() for it in th["items"]}
    blocks = {b["name"]: b["n"] for b in res["blocks"]}
    check(res["answers"] == 6 and res["targets"] == 8, "空・null は回答から除き、 対象の人数には数える")
    check(items["かき氷"] == 2 and items["海"] == 2, "項目の人数 = 当たった人の数 (1 人の中で 2 回でも 1)")
    check(by["水辺と暑さ"]["n"] == 3 and by["食べ物"]["n"] == 2, "テーマの人数 = 項目の和集合 (海 2 + 暑さ 1 で、 重なり無し 3)")
    check(blocks["夏"] == 3, "まとまりの人数 = テーマの和集合 (水辺 3 と食べ物 2 が 2 人重なり 3、 合計の 5 でも最後のテーマの 2 でもない)")
    check(by["食べ物"]["shown"] == [("かき氷", 2)], "docx の項目は 0 人を出さない")
    check([l for l, _ in by["水辺と暑さ"]["shown"]] == ["海", "暑さ"] and [l for l, _ in by["秋"]["shown"]] == ["落ち葉・紅葉", "音"],
          "docx の項目は人数の多い順、 同数は表の順")
    toks = dict(res["unmatched_tokens"])
    check(toks == {"なんとなく好きだから": 1, "どれも好き": 1},
          "どの項目にも入らない語 = 語ごとに当て、 当たらない語だけ")
    check(tokens("雪1.花見2.新学期") == ["雪", "花見", "新学期"] and tokens("「雪」、こたつ・ 音\n- 海") == ["雪", "こたつ", "音", "海"],
          "語の切り出し: 字の後の番号・括弧・読点・中黒・改行・行頭の -")
    check(res["unmatched_answers"] == 0 and build(ANS + ["何もない"], TABLE)["unmatched_answers"] == 1,
          "どの項目にも当たらない回答の人数")
    check(res["quotes"] == ["落ち葉を踏む音／紅葉", "海で泳ぐ、海の家でかき氷"], "抜粋 = key で 1 人に特定し、 改行は ／・行頭の ・ を落とす")
    e0 = err(lambda: build(ANS, {**TABLE, "quotes": {"keys": ["かき氷"]}}))
    e1 = err(lambda: build(ANS, {**TABLE, "quotes": {"keys": ["存在しない語"]}}))
    check("2 人" in e0 and "0 人" in e1 and "海で" not in e0, "抜粋の key が 0 人・2 人に当たれば止まる (回答の中身は出さない)")
    check(res["intro"] == ["回答 6 人 (対象 8 人)。 かき氷 2 人、 夏 3 人、 {未知}"],
          "差し込み: {answers} {targets} {n:} {block:} を人数に、 知らない {…} はそのまま")
    check("名前が表に無い" in err(lambda: build(ANS, {**TABLE, "intro": ["{n:無い項目}"]})), "差し込みの名前が表に無ければ止まる")
    check(res["sections"] == [("h2", "案", False), ("p", "① 太字", True), ("b", "海は 2 人", False)], "追記の節に差し込みと太字")
    bad_rx = {**TABLE, "blocks": [{"name": "x", "themes": [{"name": "y", "items": [["壊れた", "(あ"]]}]}]}
    check("壊れた" in err(lambda: build(ANS, bad_rx)), "正規表現の誤りは項目の名前つきで止まる")
    dup = {**TABLE, "blocks": TABLE["blocks"] + [{"name": "再", "themes": [{"name": "z", "items": [["海", "海"]]}]}]}
    check("重複" in err(lambda: build(ANS, dup)), "項目の名前の重複は止まる")
    check("--field" in err(lambda: build([{"ans": "雪"}], TABLE)) and build([{"ans": "雪"}, {"ans": None}], {**TABLE, "quotes": None}, "ans")["answers"] == 1,
          "object の回答は --field で欄を指定")

    with tempfile.TemporaryDirectory() as d:
        ap, tp, op = Path(d) / "a.json", Path(d) / "t.json", Path(d) / "o.docx"
        ap.write_text(json.dumps(ANS, ensure_ascii=False), encoding="utf-8")
        tp.write_text(json.dumps(TABLE, ensure_ascii=False), encoding="utf-8")
        badq = Path(d) / "bad.json"
        badq.write_text(json.dumps({**TABLE, "quotes": {"keys": ["かき氷"]}}, ensure_ascii=False), encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
            rc_bad = main(["--answers", str(ap), "--table", str(badq)])
        check(rc_bad == 2, "CLI: 表の誤りは exit 2")
        try:
            import docx  # noqa: F401
        except ImportError:
            print("SKIP docx の検査 (python-docx が無い: pip install python-docx)")
        else:
            out = io.StringIO()
            with contextlib.redirect_stdout(out), contextlib.redirect_stderr(io.StringIO()):
                rc = main(["--answers", str(ap), "--table", str(tp), "--docx", str(op)])
            check("どの項目にも入らない語 (2 種類)" in out.getvalue() and "⚠️ 0 人の項目 (docx に出ない): 食べ物 / すいか" in out.getvalue(),
                  "CLI: stdout に人数・0 人の項目・どの項目にも入らない語")
            from docx import Document
            ps = [(p.style.name, p.text, any(r.bold for r in p.runs)) for p in Document(str(op)).paragraphs]
            texts = [t for _, t, _ in ps]
            check(rc == 0 and ps[0] == ("Heading 1", "合成の季節アンケート", False), "CLI: docx を書く (見出し 1 = title)")
            check(("List Bullet", "夏 — 3 人", False) in ps and ("Heading 2", "夏（3 人）", False) in ps
                  and ("Heading 3", "水辺と暑さ（3 人）", False) in ps, "docx: まとまりの箇条書き・見出し 2・テーマの見出し 3 に人数")
            check("海 (2)、暑さ (1)" in texts, "docx: テーマの下に「項目 (人数)」 を 、 でつなぐ")
            check(("Heading 2", "抜粋 (6 人から)", False) in ps and ("List Bullet", "落ち葉を踏む音／紅葉", False) in ps,
                  "docx: 抜粋の見出しと箇条書き")
            check(ps[-3:] == [("Heading 2", "案", False), ("Normal", "① 太字", True), ("List Bullet", "海は 2 人", False)],
                  "docx: 追記の節は最後に、 太字の指定どおり")
    print("short-answer-themes selftest:", "ALL PASS" if not fails else f"FAIL {fails}")
    return 1 if fails else 0


if __name__ == "__main__":
    raise SystemExit(main())
